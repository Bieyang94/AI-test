from docx import Document
import io
from typing import List, Dict, Any, Union

from app.core.plan import reactor


def split_word_to_memory(word_source: Union[str, bytes, io.BytesIO]) -> Dict[str, Any]:
    """
    纯内存版 Word 拆分接口（支持 路径 / 二进制 / BytesIO）
    不生成任何文件，直接返回文本、表格、图片数据
    :param word_source: Word 文档路径 或 二进制bytes 或 BytesIO对象
    :return: 结构化数据（纯内存）
    """
    # ===================== 统一处理输入：支持路径、bytes、BytesIO =====================
    if isinstance(word_source, str):
        # 传入文件路径
        doc = Document(word_source)
    elif isinstance(word_source, bytes):
        # 传入二进制bytes → 转为内存流
        doc = Document(io.BytesIO(word_source))
    elif isinstance(word_source, io.BytesIO):
        # 直接传入内存流
        doc = Document(word_source)
    else:
        raise ValueError("仅支持 str路径 / bytes二进制 / BytesIO 对象")

    # ===================== 1. 提取文本（内存） =====================
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    full_text = "\n".join(paragraphs)

    # ===================== 2. 提取表格（内存） =====================
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        tables.append(table_data)

    # ===================== 3. 提取图片（内存二进制） =====================
    images = []
    for rel in doc.part._rels.values():
        if rel.reltype == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image":
            img_bytes = rel.target_part.blob
            content_type = rel.target_part.content_type
            ext = content_type.split("/")[-1].replace("jpeg", "jpg")

            images.append({
                "bytes": img_bytes,
                "ext": ext,
                "content_type": content_type
            })

    return {
        "full_text": full_text,
        "paragraphs": paragraphs,
        "tables": tables,  # 二维列表：[[行1],[行2]...]
        "images": images   # 列表：[{"bytes":..., "ext":...}, ...]
    }


def extract_text_only(word_source: Union[str, bytes, io.BytesIO]) -> str:
    """仅提取 Word 文本，不调用 LLM"""
    return split_word_to_memory(word_source)["full_text"]


def generate_title_by_model(word_source: Union[str, bytes, io.BytesIO]):
    """支持 路径 / 二进制 生成标题"""
    full_text = split_word_to_memory(word_source)["full_text"]
    extractor_title_agent = reactor.extract_title_ReactorAgent()
    requirement_docx = extractor_title_agent.execute_stream(full_text)
    return requirement_docx


def generate_csv_by_model(word_source: Union[str, bytes, io.BytesIO]):
    """支持 路径 / 二进制 生成CSV内容"""
    full_text = split_word_to_memory(word_source)["full_text"]
    test_csv_agent = reactor.ReactAgent()
    result = test_csv_agent.agent.ainvoke({
        "messages": [
            {"role": "user", "content": full_text}
        ]
    })
    return result["messages"][-1].content

# ======================== 测试代码（纯内存演示） ========================
if __name__ == "__main__":
    # 测试1：传入文件路径（原有用法，不变）
    TEST_FILE = r"C:\Users\EDY\Desktop\掘金2.0产品需求文档 .docx"
    # result = generate_csv_by_model(TEST_FILE)

    # 测试2：传入二进制bytes（新增功能）
    with open(TEST_FILE, "rb") as f:
        word_bytes = f.read()  # 读取为二进制
    # result = generate_csv_by_model(word_bytes)  # 直接传bytes
    result = generate_title_by_model(word_bytes)

    # print(result)